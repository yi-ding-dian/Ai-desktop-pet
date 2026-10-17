"""DOCX 解析单测：段落顺序、表格 Markdown/TSV 降级、合并单元格、老式 .doc 拒绝。"""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from docx import Document

from backend.modules.files import FileLimits, FileRef, FileRejected, read_attachments
from backend.modules.files.docx_reader import MD_MAX_COLS, MD_MAX_ROWS


def _limits(**overrides: int) -> FileLimits:
    values = {"max_chars": 20_000, "max_bytes": 1024 * 1024, "max_per_turn": 3}
    values.update(overrides)
    return FileLimits(**values)


def _save_docx(path: Path, build) -> Path:
    document = Document()
    build(document)
    document.save(path)
    return path


def test_paragraphs_and_small_table_in_document_order(tmp_path: Path) -> None:
    def build(document: Document) -> None:
        document.add_paragraph("开头段落")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "姓名"
        table.cell(0, 1).text = "年龄"
        table.cell(1, 0).text = "张三"
        table.cell(1, 1).text = "28"
        document.add_paragraph("结尾段落")

    path = _save_docx(tmp_path / "report.docx", build)
    text = read_attachments([FileRef(str(path))], _limits())[0].text

    assert text.index("开头段落") < text.index("表 1（2 行 × 2 列）") < text.index("结尾段落")
    assert "| 列1 | 列2 |" in text
    assert "| 姓名 | 年龄 |" in text
    assert "| 张三 | 28 |" in text


def test_wide_table_degrades_to_tsv(tmp_path: Path) -> None:
    cols = MD_MAX_COLS + 1

    def build(document: Document) -> None:
        table = document.add_table(rows=2, cols=cols)
        for row in range(2):
            for col in range(cols):
                table.cell(row, col).text = f"r{row}c{col}"

    path = _save_docx(tmp_path / "wide.docx", build)
    text = read_attachments([FileRef(str(path))], _limits())[0].text

    assert f"表 1（2 行 × {cols} 列）" in text
    assert "（含合并单元格）" not in text  # 普通表不能被误判成合并
    assert "\t" in text
    assert "| 列1 |" not in text  # 降级后没有 Markdown 表头
    # 所有单元格都必须保留（曾因 tc 代理 id 复用导致整行被吞）
    for row in range(2):
        for col in range(cols):
            assert f"r{row}c{col}" in text


def test_long_table_degrades_to_tsv(tmp_path: Path) -> None:
    rows = MD_MAX_ROWS + 1

    def build(document: Document) -> None:
        table = document.add_table(rows=rows, cols=2)
        for row in range(rows):
            table.cell(row, 0).text = f"k{row}"
            table.cell(row, 1).text = f"v{row}"

    path = _save_docx(tmp_path / "long.docx", build)
    text = read_attachments([FileRef(str(path))], _limits())[0].text

    assert f"表 1（{rows} 行 × 2 列）" in text
    assert "\t" in text


def test_merged_cells_marked_and_value_written_once(tmp_path: Path) -> None:
    def build(document: Document) -> None:
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1))
        table.cell(0, 0).text = "合并标题"
        table.cell(1, 0).text = "左"
        table.cell(1, 1).text = "右"

    path = _save_docx(tmp_path / "merged.docx", build)
    text = read_attachments([FileRef(str(path))], _limits())[0].text

    assert "（含合并单元格）" in text
    assert text.count("合并标题") == 1
    assert "| 左 | 右 |" in text


def test_legacy_doc_rejected_with_guidance(tmp_path: Path) -> None:
    path = tmp_path / "old.doc"
    path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 32)

    with pytest.raises(FileRejected) as exc:
        read_attachments([FileRef(str(path))], _limits())
    assert exc.value.code == "DOC_LEGACY"
    assert "另存为 .docx" in exc.value.message


def test_zip_without_docx_body_rejected(tmp_path: Path) -> None:
    path = tmp_path / "fake.docx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("hello.txt", "not a word document")

    with pytest.raises(FileRejected) as exc:
        read_attachments([FileRef(str(path))], _limits())
    assert exc.value.code == "DOCX_INVALID"


def test_docx_without_text_rejected(tmp_path: Path) -> None:
    path = _save_docx(tmp_path / "empty.docx", lambda document: None)

    with pytest.raises(FileRejected) as exc:
        read_attachments([FileRef(str(path))], _limits())
    assert exc.value.code == "DOCX_NO_TEXT"


def test_docx_text_counts_toward_char_limit(tmp_path: Path) -> None:
    def build(document: Document) -> None:
        document.add_paragraph("字" * 100)

    path = _save_docx(tmp_path / "long-text.docx", build)

    with pytest.raises(FileRejected) as exc:
        read_attachments([FileRef(str(path))], _limits(max_chars=50))
    assert exc.value.code == "FILE_TOO_LONG"
    assert "50" in exc.value.message
