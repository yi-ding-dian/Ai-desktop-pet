"""Excel 解析单测：多 sheet/隐藏 sheet、TSV、正式表格、合并单元格、公式兜底。"""

from __future__ import annotations

import datetime as dt
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.table import Table

from backend.modules.files import FileLimits, FileRef, FileRejected, read_attachments
from backend.modules.files import excel_reader


def _limits(**overrides: int) -> FileLimits:
    values = {"max_chars": 20_000, "max_bytes": 5 * 1024 * 1024, "max_per_turn": 3}
    values.update(overrides)
    return FileLimits(**values)


def _save_workbook(path: Path, build) -> Path:
    workbook = Workbook()
    build(workbook)
    workbook.save(path)
    return path


def _read_text(path: Path, **limits: int) -> str:
    return read_attachments([FileRef(str(path))], _limits(**limits))[0].text


def test_sheet_values_dates_and_percent_as_tsv(tmp_path: Path) -> None:
    def build(workbook: Workbook) -> None:
        ws = workbook.active
        ws.title = "销售"
        ws.append(["地区", "金额", "日期", "进度"])
        ws.append(["华东", 1200, dt.datetime(2026, 1, 1), 0.8])
        ws["D2"].number_format = "0%"

    path = _save_workbook(tmp_path / "sales.xlsx", build)
    text = _read_text(path)

    assert "工作表：销售（2 行 × 4 列）" in text
    assert "地区\t金额\t日期\t进度" in text
    assert "华东\t1200\t2026-01-01\t80%" in text


def test_multiple_sheets_and_hidden_sheet_skipped(tmp_path: Path) -> None:
    def build(workbook: Workbook) -> None:
        first = workbook.active
        first.title = "一表"
        first["A1"] = "one"
        second = workbook.create_sheet("二表")
        second["A1"] = "two"
        hidden = workbook.create_sheet("隐藏表")
        hidden["A1"] = "secret"
        hidden.sheet_state = "hidden"

    path = _save_workbook(tmp_path / "sheets.xlsx", build)
    text = _read_text(path)

    assert "工作表：一表" in text
    assert "工作表：二表" in text
    assert "隐藏表" not in text
    assert "已跳过 1 个隐藏工作表" in text


def test_merged_cells_manifest_and_blank_grid(tmp_path: Path) -> None:
    def build(workbook: Workbook) -> None:
        ws = workbook.active
        ws.title = "合并"
        ws.merge_cells("A1:B1")
        ws["A1"] = "合并标题"
        ws.merge_cells("A2:A3")
        ws["A2"] = "华东"
        ws["B2"] = "A"
        ws["C2"] = 100
        ws["B3"] = "B"
        ws["C3"] = 120

    path = _save_workbook(tmp_path / "merged.xlsx", build)
    text = _read_text(path)

    assert "（含合并单元格）" in text
    assert "合并单元格：A1:B1=合并标题；A2:A3=华东" in text
    assert "华东\tA\t100" in text
    # A3 是被合并掉的空位，第二行以空字段开头
    assert "\tB\t120" in text


def test_formal_tables_split_and_other_cells(tmp_path: Path) -> None:
    def build(workbook: Workbook) -> None:
        ws = workbook.active
        ws.title = "Sheet1"
        ws.append(["Region", "Amount", "Qty"])
        ws.append(["East", 100, 2])
        ws.append(["South", 80, 1])
        ws.add_table(Table(displayName="Sales", ref="A1:C3"))
        ws["E1"] = "Metric"
        ws["F1"] = "Value"
        ws["E2"] = "Total"
        ws["F2"] = 180
        ws.add_table(Table(displayName="Summary", ref="E1:F2"))
        ws["A5"] = "备注：仅测试"

    path = _save_workbook(tmp_path / "tables.xlsx", build)
    text = _read_text(path)

    assert "表格：Sales（A1:C3，3 行 × 3 列）" in text
    assert "Region\tAmount\tQty" in text
    assert "表格：Summary（E1:F2，2 行 × 2 列）" in text
    assert "Metric\tValue" in text
    assert "其他单元格：A5=备注：仅测试" in text


def test_formula_without_cached_value_falls_back_to_formula(tmp_path: Path) -> None:
    def build(workbook: Workbook) -> None:
        ws = workbook.active
        ws.title = "公式"
        ws["A1"] = 1
        ws["B1"] = 2
        ws["C1"] = "=A1+B1"

    path = _save_workbook(tmp_path / "formula.xlsx", build)
    text = _read_text(path)

    assert "=A1+B1（未计算，无缓存结果）" in text
    assert "1 个公式单元格没有缓存结果" in text


def test_char_limit_rejects_whole_workbook(tmp_path: Path) -> None:
    def build(workbook: Workbook) -> None:
        ws = workbook.active
        ws["A1"] = "字" * 100

    path = _save_workbook(tmp_path / "long.xlsx", build)

    with pytest.raises(FileRejected) as exc:
        _read_text(path, max_chars=50)
    assert exc.value.code == "FILE_TOO_LONG"
    assert "50" in exc.value.message


def test_cell_count_cap_rejects_large_used_range(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # 用较小的上限测逻辑，避免真的构造 10 万行文件拖慢测试
    monkeypatch.setattr(excel_reader, "MAX_CELLS_PER_SHEET", 5)

    def build(workbook: Workbook) -> None:
        ws = workbook.active
        ws.append(["a", "b", "c"])
        ws.append(["d", "e", "f"])

    path = _save_workbook(tmp_path / "huge.xlsx", build)

    with pytest.raises(FileRejected) as exc:
        _read_text(path)
    assert exc.value.code == "XLSX_TOO_LARGE"


def test_legacy_xls_rejected_with_guidance(tmp_path: Path) -> None:
    path = tmp_path / "old.xls"
    path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 32)

    with pytest.raises(FileRejected) as exc:
        _read_text(path)
    assert exc.value.code == "XLS_LEGACY"
    assert "另存为 .xlsx" in exc.value.message


def test_encrypted_xlsx_ole_container_gets_clear_message(tmp_path: Path) -> None:
    path = tmp_path / "locked.xlsx"
    path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 32)

    with pytest.raises(FileRejected) as exc:
        _read_text(path)
    assert exc.value.code == "XLSX_ENCRYPTED"
    assert "加密" in exc.value.message


def test_zip_without_workbook_rejected(tmp_path: Path) -> None:
    path = tmp_path / "fake.xlsx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("hello.txt", "not excel")

    with pytest.raises(FileRejected) as exc:
        _read_text(path)
    assert exc.value.code == "XLSX_INVALID"


def test_empty_workbook_rejected(tmp_path: Path) -> None:
    path = _save_workbook(tmp_path / "empty.xlsx", lambda workbook: None)

    with pytest.raises(FileRejected) as exc:
        _read_text(path)
    assert exc.value.code == "XLSX_NO_DATA"
